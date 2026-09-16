from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from app.database import engine, Base
from app.models import *

from app.api.config import router as config_router
from app.api.inventory import router as inventory_router
from app.api.production import router as production_router
from app.api.sales import router as sales_router
from app.api.accounting import router as accounting_router
from app.api.auth import router as auth_router

# Sincronizar esquemas de base de datos SQLite
Base.metadata.create_all(bind=engine)

# Migración segura de columnas adicionales en SQLite
try:
    from sqlalchemy import text
    with engine.connect() as conn:
        res = conn.execute(text("PRAGMA table_info(raw_materials)"))
        cols = [r[1] for r in res.fetchall()]
        if cols and "created_at" not in cols:
            conn.execute(text("ALTER TABLE raw_materials ADD COLUMN created_at DATETIME"))
            conn.execute(text("UPDATE raw_materials SET created_at = updated_at WHERE created_at IS NULL"))
            conn.commit()

        if cols and "article_code" not in cols:
            conn.execute(text("ALTER TABLE raw_materials ADD COLUMN article_code TEXT"))
            conn.commit()

        res_supp = conn.execute(text("PRAGMA table_info(additional_supplies)"))
        cols_supp = [r[1] for r in res_supp.fetchall()]
        if cols_supp and "created_at" not in cols_supp:
            conn.execute(text("ALTER TABLE additional_supplies ADD COLUMN created_at DATETIME"))
            conn.execute(text("UPDATE additional_supplies SET created_at = updated_at WHERE created_at IS NULL"))
            conn.commit()
except Exception as e:
    pass

# Auto-asignar código de artículo oficial a cualquier material existente que tenga '-' o nulo
# e importar catálogo oficial de 60 bobinas si la base de datos tiene menos de 10 registros
try:
    from app.database import SessionLocal, PROJECT_ROOT
    from app.models.inventory import RawMaterial
    from app.utils.article_codes import generate_article_code
    import os
    db_init = SessionLocal()
    mats_without_code = db_init.query(RawMaterial).filter(
        (RawMaterial.article_code == None) | (RawMaterial.article_code == "") | (RawMaterial.article_code == "-")
    ).all()
    for m in mats_without_code:
        m.article_code = generate_article_code(m.material_type or "PETG", m.color or "Blanco", db=db_init)
    if mats_without_code:
        db_init.commit()

    # Si la base de datos tiene menos de 10 materiales, auto-sincronizar el catálogo oficial de bobinas del Excel
    if db_init.query(RawMaterial).count() < 10:
        import openpyxl
        excel_path = os.path.join(PROJECT_ROOT, "Sistema_Integral_Produccion_ prisma lab(Recuperado automáticamente).xlsm")
        if os.path.exists(excel_path):
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            sheet_name = "Inventario_Materiales_(2)" if "Inventario_Materiales_(2)" in wb.sheetnames else "Inventario_Materiales"
            sheet = wb[sheet_name]
            start_row = 7 if sheet_name == "Inventario_Materiales_(2)" else 2
            for r in range(start_row, sheet.max_row + 1):
                raw_code = sheet.cell(row=r, column=1).value if sheet_name == "Inventario_Materiales_(2)" else None
                raw_name = sheet.cell(row=r, column=2 if sheet_name == "Inventario_Materiales_(2)" else 1).value
                if not raw_name:
                    continue
                name_str = str(raw_name).strip()
                if not name_str:
                    continue
                
                code_str = str(raw_code).strip() if raw_code else None
                if code_str and db_init.query(RawMaterial).filter(RawMaterial.article_code == code_str).first():
                    continue
                if db_init.query(RawMaterial).filter(RawMaterial.name == name_str).first():
                    continue

                color_val = str(sheet.cell(row=r, column=3 if sheet_name == "Inventario_Materiales_(2)" else 2).value or "Estándar").strip()
                mtype_val = str(sheet.cell(row=r, column=4 if sheet_name == "Inventario_Materiales_(2)" else 3).value or "PLA").strip()
                try:
                    stock_init_val = float(sheet.cell(row=r, column=7 if sheet_name == "Inventario_Materiales_(2)" else 4).value or 1000.0)
                except Exception:
                    stock_init_val = 1000.0
                try:
                    outgoing_val = float(sheet.cell(row=r, column=8 if sheet_name == "Inventario_Materiales_(2)" else 5).value or 0.0)
                except Exception:
                    outgoing_val = 0.0
                try:
                    cost_g_val = float(sheet.cell(row=r, column=10 if sheet_name == "Inventario_Materiales_(2)" else 8).value or 65.0)
                except Exception:
                    cost_g_val = 65.0
                if cost_g_val <= 0:
                    cost_g_val = 65.0
                
                notes_val = str(sheet.cell(row=r, column=14 if sheet_name == "Inventario_Materiales_(2)" else 10).value or "").strip()
                prov_val = str(sheet.cell(row=r, column=5).value or "").strip() if sheet_name == "Inventario_Materiales_(2)" else ""
                if prov_val:
                    notes_val = f"Proveedor: {prov_val}. {notes_val}".strip()

                mat = RawMaterial(
                    article_code=code_str if code_str else generate_article_code(mtype_val, color_val, db=db_init),
                    name=name_str,
                    color=color_val.capitalize() if color_val else "Estándar",
                    material_type=mtype_val.upper() if mtype_val else "PLA",
                    initial_stock_g=stock_init_val,
                    outgoing_stock_g=outgoing_val,
                    current_stock_g=stock_init_val - outgoing_val,
                    cost_per_g=cost_g_val,
                    notes=notes_val,
                    min_stock_alert_g=200.0
                )
                db_init.add(mat)
            db_init.commit()

    # 3. Auto-sincronizar y reclasificar asientos de Flujo de Caja (Equipos a Activos Fijos 152005, Capital a Patrimonio 311505)
    try:
        from datetime import datetime
        from app.models.accounting import CashFlowRecord, JournalEntry
        from app.api.accounting import map_cashflow_to_puc

        cash_records = db_init.query(CashFlowRecord).all()
        for cf in cash_records:
            is_income = (cf.income or 0) > 0
            amount = cf.income if is_income else cf.credit
            if not amount or amount <= 0:
                continue

            target_puc, target_name = map_cashflow_to_puc(cf.category, cf.description, is_income)

            search_prefix = f"Flujo de Caja ({'Ingreso' if is_income else 'Egreso'}): {cf.description}"
            matching_entries = db_init.query(JournalEntry).filter(JournalEntry.description == search_prefix).all()

            if matching_entries:
                for e in matching_entries:
                    if is_income:
                        if e.credit > 0 and e.puc_code != target_puc:
                            e.puc_code = target_puc
                            e.account_name = target_name
                    else:
                        if e.debit > 0 and e.puc_code != target_puc:
                            e.puc_code = target_puc
                            e.account_name = target_name
            else:
                last_je = db_init.query(JournalEntry).order_by(JournalEntry.entry_number.desc()).first()
                next_num = (last_je.entry_number if last_je else 0) + 1
                entry_dt = cf.record_date or datetime.utcnow()
                if is_income:
                    db_init.add(JournalEntry(
                        entry_number=next_num,
                        entry_date=entry_dt,
                        puc_code="110505",
                        account_name="Caja General",
                        description=search_prefix,
                        debit=round(amount, 2),
                        credit=0.0
                    ))
                    db_init.add(JournalEntry(
                        entry_number=next_num,
                        entry_date=entry_dt,
                        puc_code=target_puc,
                        account_name=target_name,
                        description=search_prefix,
                        debit=0.0,
                        credit=round(amount, 2)
                    ))
                else:
                    db_init.add(JournalEntry(
                        entry_number=next_num,
                        entry_date=entry_dt,
                        puc_code=target_puc,
                        account_name=target_name,
                        description=search_prefix,
                        debit=round(amount, 2),
                        credit=0.0
                    ))
                    db_init.add(JournalEntry(
                        entry_number=next_num,
                        entry_date=entry_dt,
                        puc_code="110505",
                        account_name="Caja General",
                        description=search_prefix,
                        debit=0.0,
                        credit=round(amount, 2)
                    ))
        db_init.commit()
    except Exception as e:
        db_init.rollback()

    # 4. Auto-crear usuario administrador inicial si no existen usuarios
    try:
        from app.models.auth import User
        from app.core.security import hash_password
        admin_user = db_init.query(User).filter(User.username == "admin").first()
        if not admin_user and db_init.query(User).count() == 0:
            initial_admin = User(
                username="admin",
                hashed_password=hash_password("admin123"),
                full_name="Administrador Prisma Lab",
                role="ADMIN",
                is_active=True
            )
            db_init.add(initial_admin)
            db_init.commit()
            print("[OK] Usuario inicial 'admin' creado exitosamente (Clave: admin123).")
    except Exception as e:
        db_init.rollback()

    db_init.close()
except Exception as e:
    pass

app = FastAPI(
    title="Prisma Lab ERP - API Backend",
    description="Sistema Integral de Producción, Cotización y Contabilidad 3D (Migración Excel a Web)",
    version="1.0.0"
)

# Habilitar CORS para frontend local React Vite
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Registrar Routers API con prefijo /api
app.include_router(config_router, prefix="/api")
app.include_router(inventory_router, prefix="/api")
app.include_router(production_router, prefix="/api")
app.include_router(sales_router, prefix="/api")
app.include_router(accounting_router, prefix="/api")
app.include_router(auth_router, prefix="/api")

@app.get("/")
def read_root():
    return {
        "status": "online",
        "app": "Prisma Lab ERP & 3D Production System",
        "version": "1.0.0",
        "docs": "/docs"
    }

@app.get("/api/health")
def health_check():
    from app.database import SessionLocal, DEFAULT_DB_PATH
    from app.models.inventory import RawMaterial
    from app.models.accounting import JournalEntry
    db = SessionLocal()
    try:
        mat_count = db.query(RawMaterial).count()
        last_mat = db.query(RawMaterial).order_by(RawMaterial.id.desc()).first()
        entry_count = db.query(JournalEntry).count()
        last_entry = db.query(JournalEntry).order_by(JournalEntry.id.desc()).first()
        return {
            "status": "healthy",
            "backend_version": "v1.4",
            "database_file": DEFAULT_DB_PATH,
            "materials_count": mat_count,
            "last_material": {
                "id": last_mat.id,
                "code": last_mat.article_code,
                "name": last_mat.name,
                "created_at": str(last_mat.created_at)
            } if last_mat else None,
            "journal_entries_count": entry_count,
            "last_journal_entry": {
                "entry_number": last_entry.entry_number,
                "puc": last_entry.puc_code,
                "date": str(last_entry.entry_date),
                "description": last_entry.description
            } if last_entry else None
        }
    finally:
        db.close()
